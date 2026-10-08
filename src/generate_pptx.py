import os
import json
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

# Dynamic Path Resolution
script_dir = os.path.dirname(os.path.abspath(__file__))
project_root = os.path.dirname(script_dir)

fig_dir = os.path.join(project_root, 'visualizations')
metrics_file = os.path.join(project_root, 'metrics', 'summary_metrics.json')
output_dir = os.path.join(project_root, 'presentations')
os.makedirs(output_dir, exist_ok=True)

# Load summary metrics
with open(metrics_file, 'r') as f:
    metrics = json.load(f)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Color Palette
COLOR_BG_DARK = RGBColor(15, 23, 42)      # Slate 900
COLOR_PRIMARY = RGBColor(30, 58, 138)     # Blue 900
COLOR_ACCENT_BLUE = RGBColor(37, 99, 235) # Blue 600
COLOR_ACCENT_TEAL = RGBColor(13, 148, 136)# Teal 600
COLOR_TEXT_MAIN = RGBColor(15, 23, 42)    # Slate 900
COLOR_TEXT_MUTED = RGBColor(100, 116, 139)# Slate 500
COLOR_CARD_BG = RGBColor(248, 250, 252)   # Slate 50
COLOR_BORDER = RGBColor(226, 232, 240)    # Slate 200
COLOR_SUCCESS = RGBColor(16, 185, 129)    # Green 500
COLOR_DANGER = RGBColor(239, 68, 68)      # Red 500
COLOR_WHITE = RGBColor(255, 255, 255)

def add_header(slide, title_text, category_text="FINTECH OPERATIONS & DATA ANALYTICS"):
    cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.35))
    tf_c = cat_box.text_frame
    tf_c.word_wrap = True
    tf_c.margin_left = tf_c.margin_right = tf_c.margin_top = tf_c.margin_bottom = 0
    p_c = tf_c.paragraphs[0]
    p_c.text = category_text.upper()
    p_c.font.size = Pt(10)
    p_c.font.bold = True
    p_c.font.color.rgb = COLOR_ACCENT_BLUE

    t_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.75), Inches(11.7), Inches(0.6))
    tf_t = t_box.text_frame
    tf_t.word_wrap = True
    tf_t.margin_left = tf_t.margin_right = tf_t.margin_top = tf_t.margin_bottom = 0
    p_t = tf_t.paragraphs[0]
    p_t.text = title_text
    p_t.font.size = Pt(22)
    p_t.font.bold = True
    p_t.font.color.rgb = COLOR_TEXT_MAIN

def add_card(slide, left, top, width, height, bg_color=COLOR_CARD_BG, border_color=COLOR_BORDER):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = bg_color
    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1)
    else:
        shape.line.fill.background()
    return shape

# ==============================================================================
# SLIDE 1: TITLE / COVER
# ==============================================================================
blank_layout = prs.slide_layouts[6]
s1 = prs.slides.add_slide(blank_layout)

bg = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
bg.fill.solid()
bg.fill.fore_color.rgb = COLOR_BG_DARK
bg.line.fill.background()

tb = s1.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(11.333), Inches(3.8))
tf = tb.text_frame
tf.word_wrap = True

p0 = tf.paragraphs[0]
p0.text = "DATA ANALYTICS CASE STUDY"
p0.font.size = Pt(14)
p0.font.bold = True
p0.font.color.rgb = RGBColor(56, 189, 248)
p0.space_after = Pt(14)

p1 = tf.add_paragraph()
p1.text = "Fintech & PPOB Transaction Operations Analytics"
p1.font.size = Pt(32)
p1.font.bold = True
p1.font.color.rgb = COLOR_WHITE
p1.space_after = Pt(10)

p2 = tf.add_paragraph()
p2.text = "Volume Concentration, Biller Failure Bottlenecks, Peak Load Patterns & Vendor Risk"
p2.font.size = Pt(18)
p2.font.color.rgb = RGBColor(203, 213, 225)
p2.space_after = Pt(30)

p3 = tf.add_paragraph()
p3.text = "Dataset: 316,376 masked transactions (1, 5 and 6 Aug 2026) | Tech Stack: SQL, Python (pandas/Seaborn), Power BI | Author: Novaldi Ramadhan Waluyo"
p3.font.size = Pt(12)
p3.font.color.rgb = RGBColor(148, 163, 184)

# ==============================================================================
# SLIDE 2: EXECUTIVE SUMMARY & SCORECARD
# ==============================================================================
s2 = prs.slides.add_slide(blank_layout)
add_header(s2, "Executive Summary: Operational Performance")

kpis = [
    ("Total Volume", f"{metrics['total_transactions']:,}", "Analyzed across 3 operational days", COLOR_PRIMARY),
    ("Success Rate", f"{metrics['success_rate_pct']}%", f"{metrics['total_success']:,} settled transactions", COLOR_SUCCESS),
    ("Failure Rate", f"{metrics['failure_rate_pct']}%", f"{metrics['total_failed']:,} failed transactions", COLOR_DANGER),
    ("Failure Concentration", "60.9%", "19,585 failures from 2 billers", RGBColor(217, 119, 6))
]

card_w = Inches(2.7)
card_h = Inches(1.4)
card_gap = Inches(0.3)
left_start = Inches(0.8)

for i, (title, val, sub, col) in enumerate(kpis):
    c_left = left_start + i * (card_w + card_gap)
    add_card(s2, c_left, Inches(1.5), card_w, card_h)
    
    tb = s2.shapes.add_textbox(c_left + Inches(0.15), Inches(1.6), card_w - Inches(0.3), card_h - Inches(0.2))
    tf = tb.text_frame
    tf.word_wrap = True
    
    p0 = tf.paragraphs[0]
    p0.text = title.upper()
    p0.font.size = Pt(10)
    p0.font.bold = True
    p0.font.color.rgb = COLOR_TEXT_MUTED
    
    p1 = tf.add_paragraph()
    p1.text = val
    p1.font.size = Pt(22)
    p1.font.bold = True
    p1.font.color.rgb = col
    
    p2 = tf.add_paragraph()
    p2.text = sub
    p2.font.size = Pt(9)
    p2.font.color.rgb = COLOR_TEXT_MUTED

s2.shapes.add_picture(os.path.join(fig_dir, '01_overall_status_distribution.png'), Inches(1.0), Inches(3.1), height=Inches(3.9))

add_card(s2, Inches(6.6), Inches(3.1), Inches(5.9), Inches(3.9))
tb_right = s2.shapes.add_textbox(Inches(6.8), Inches(3.2), Inches(5.5), Inches(3.7))
tf_r = tb_right.text_frame
tf_r.word_wrap = True

p_r0 = tf_r.paragraphs[0]
p_r0.text = "Key Operational Highlights & Problem Statement"
p_r0.font.size = Pt(14)
p_r0.font.bold = True
p_r0.font.color.rgb = COLOR_PRIMARY
p_r0.space_after = Pt(10)

findings = [
    ("Volume Baseline", "316.4k transactions over three days with an overall 89.84% success rate."),
    ("Failure Concentration", "60.9% of all failures (19,585) come from two billers, Biller_29 and Biller_26, which handle only 19.6% of traffic."),
    ("High-Volume Product Risk", "TNP23, the #3 SKU by volume (22.3k tx), fails 22.75% of the time."),
    ("Broken Denominations", "XDF1000 (6.2k tx) fails 58.73% of the time, XDF2000 53.47% and XDF3000 48.28%, pointing to an upstream catalogue or denomination problem."),
    ("Data Quality", "1,923 transactions (0.6%) used lower-case product codes such as xdg1. They failed 34.9% of the time vs 10.0% for standard codes. Codes were upper-cased before analysis.")
]

for title, desc in findings:
    p = tf_r.add_paragraph()
    p.text = f"• {title}: {desc}"
    p.font.size = Pt(10.5)
    p.font.color.rgb = COLOR_TEXT_MAIN
    p.space_after = Pt(6)

# ==============================================================================
# SLIDE 3: PRODUCT POPULARITY & PARETO CONCENTRATION
# ==============================================================================
s3 = prs.slides.add_slide(blank_layout)
add_header(s3, "Product Volume Distribution: Pareto Concentration")

s3.shapes.add_picture(os.path.join(fig_dir, '06_pareto_volume_concentration.png'), Inches(0.8), Inches(1.5), width=Inches(6.5))

add_card(s3, Inches(7.5), Inches(1.5), Inches(5.0), Inches(5.5))
tb_p3 = s3.shapes.add_textbox(Inches(7.7), Inches(1.7), Inches(4.6), Inches(5.1))
tf_p3 = tb_p3.text_frame
tf_p3.word_wrap = True

p = tf_p3.paragraphs[0]
p.text = "Top Volume Drivers"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = COLOR_PRIMARY
p.space_after = Pt(12)

points = [
    ("Heavy Pareto Skew", "The top 5 SKUs carry 57.3% of volume (181.3k tx). It takes 20 SKUs to pass 80% (81.4%)."),
    ("Leading SKU (SB20)", "79,184 transactions (25.0% of volume) with a 94.70% success rate."),
    ("Runner-up (XDG1)", "41,492 transactions (13.1% of volume) with a 96.47% success rate."),
    ("Volume Vulnerability (TNP23)", "Ranked #3 in volume (22,334 tx) but plagued by a 22.75% failure rate (5,082 failed tx)."),
    ("Implication", "The top 3 SKUs carry 45.2% of volume, so close monitoring of just three products protects almost half of all transactions.")
]

for t, d in points:
    p_item = tf_p3.add_paragraph()
    p_item.text = f"• {t}: {d}"
    p_item.font.size = Pt(11)
    p_item.font.color.rgb = COLOR_TEXT_MAIN
    p_item.space_after = Pt(8)

# ==============================================================================
# SLIDE 4: OPERATIONAL FAILURE HOTSPOTS & PRODUCT SLA
# ==============================================================================
s4 = prs.slides.add_slide(blank_layout)
add_header(s4, "Operational Hotspots: Products with Critical Failure Rates")

s4.shapes.add_picture(os.path.join(fig_dir, '03_critical_high_failure_products.png'), Inches(0.8), Inches(1.5), width=Inches(6.2))
s4.shapes.add_picture(os.path.join(fig_dir, '02_top10_products_volume_and_failure_rate.png'), Inches(7.1), Inches(1.5), width=Inches(5.4))

add_card(s4, Inches(0.8), Inches(5.6), Inches(11.7), Inches(1.4), bg_color=RGBColor(254, 242, 242), border_color=COLOR_DANGER)
tb_c4 = s4.shapes.add_textbox(Inches(1.0), Inches(5.7), Inches(11.3), Inches(1.2))
tf_c4 = tb_c4.text_frame
tf_c4.word_wrap = True

p_c4_0 = tf_c4.paragraphs[0]
p_c4_0.text = "CRITICAL DIAGNOSTIC FINDINGS:"
p_c4_0.font.size = Pt(11)
p_c4_0.font.bold = True
p_c4_0.font.color.rgb = COLOR_DANGER
p_c4_0.space_after = Pt(4)

p_c4_1 = tf_c4.add_paragraph()
p_c4_1.text = "1. XDF series (Biller_29): XDF1000, XDF2000 and XDF3000 fail 58.73%, 53.47% and 48.28% of the time, which points to a denomination mismatch or a stock/balance problem upstream.\n2. TNP series (Biller_26): TNP13 (37.35%) and TNP23 (22.75%) are high-demand products with high failure. No other biller carries them in this data, so a backup supplier is needed before any failover."
p_c4_1.font.size = Pt(10.5)
p_c4_1.font.color.rgb = RGBColor(127, 29, 29)

# ==============================================================================
# SLIDE 5: BILLER PERFORMANCE & VENDOR SLA VULNERABILITY
# ==============================================================================
s5 = prs.slides.add_slide(blank_layout)
add_header(s5, "Biller Performance: Volume vs Failure Concentration")

s5.shapes.add_picture(os.path.join(fig_dir, '04_biller_performance_and_sla.png'), Inches(0.8), Inches(1.5), width=Inches(8.0))

add_card(s5, Inches(9.0), Inches(1.5), Inches(3.5), Inches(5.5))
tb_s5 = s5.shapes.add_textbox(Inches(9.15), Inches(1.65), Inches(3.2), Inches(5.2))
tf_s5 = tb_s5.text_frame
tf_s5.word_wrap = True

p = tf_s5.paragraphs[0]
p.text = "Biller Insights"
p.font.size = Pt(13)
p.font.bold = True
p.font.color.rgb = COLOR_PRIMARY
p.space_after = Pt(10)

b_insights = [
    ("Biller_27 Backbone", "Handles 50.4% of traffic (159.4k tx) at a 94.3% success rate."),
    ("Biller_29 Alert", "60.6% success rate (39.38% failure rate); 31.4% of all failures."),
    ("Biller_26 Bottleneck", "73.96% success rate (26.04% failure rate); 29.5% of all failures."),
    ("Vendor Risk", "Together they cause 60.9% of failures on only 19.6% of volume."),
    ("No Backup Route", "Their SKUs are not served by any other biller in the data (2 of 19,585 failures), so rerouting first needs a new supplier.")
]

for t, d in b_insights:
    p_i = tf_s5.add_paragraph()
    p_i.text = f"• {t}: {d}"
    p_i.font.size = Pt(9.5)
    p_i.font.color.rgb = COLOR_TEXT_MAIN
    p_i.space_after = Pt(6)

# ==============================================================================
# SLIDE 6: 24-HOUR TEMPORAL DYNAMICS & TRAFFIC SPIKES
# ==============================================================================
s6 = prs.slides.add_slide(blank_layout)
add_header(s6, "Temporal Load Analysis: 24-Hour Traffic Curve & Failure Patterns")

s6.shapes.add_picture(os.path.join(fig_dir, '05_hourly_traffic_load_and_failure_trend.png'), Inches(0.8), Inches(1.5), width=Inches(7.5))

add_card(s6, Inches(8.5), Inches(1.5), Inches(4.0), Inches(5.5))
tb_s6 = s6.shapes.add_textbox(Inches(8.7), Inches(1.7), Inches(3.6), Inches(5.1))
tf_s6 = tb_s6.text_frame
tf_s6.word_wrap = True

p = tf_s6.paragraphs[0]
p.text = "Traffic & Load Patterns"
p.font.size = Pt(13)
p.font.bold = True
p.font.color.rgb = COLOR_PRIMARY
p.space_after = Pt(10)

t_points = [
    ("Evening Peak (17:00 - 19:00)", "18:00 WIB is the busiest hour: 24,787 transactions over the three days (~8.3k per day). 17:00-19:00 carries 21.3% of volume."),
    ("Morning Load (07:00 - 09:00)", "About 6.7k transactions per hour per day."),
    ("Off-Peak (01:00 - 04:00)", "02:00-03:00 WIB is the quietest window (~530 transactions per hour per day), suitable for batch reconciliation."),
    ("Load vs Failure", "Hourly failure rate ranges from 7.0% to 13.4% and does not rise with volume (18:00 is 9.1%), so failures are driven by billers and SKUs, not load.")
]

for t, d in t_points:
    p_i = tf_s6.add_paragraph()
    p_i.text = f"• {t}: {d}"
    p_i.font.size = Pt(10)
    p_i.font.color.rgb = COLOR_TEXT_MAIN
    p_i.space_after = Pt(6)

# ==============================================================================
# SLIDE 7: B2B PARTNER IMPACT & CUSTOMER SEGMENTATION
# ==============================================================================
s7 = prs.slides.add_slide(blank_layout)
add_header(s7, "B2B Partner Ecosystem & Customer Experience Risk")

c_w = Inches(3.7)
c_h = Inches(5.3)

add_card(s7, Inches(0.8), Inches(1.6), c_w, c_h)
tb1 = s7.shapes.add_textbox(Inches(1.0), Inches(1.8), c_w - Inches(0.4), c_h - Inches(0.4))
tf1 = tb1.text_frame
tf1.word_wrap = True
p = tf1.paragraphs[0]
p.text = "B2B Partner Exposure"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = COLOR_PRIMARY
p.space_after = Pt(8)

p1 = tf1.add_paragraph()
p1.text = "• 46 Unique B2B Partners integrated into the gateway.\n• Top 5 Partners account for 66.0% (208.9k tx) of total system volume.\n• Partner_07 is the largest client (79,974 tx), followed by Partner_12 (55,304 tx).\n• Risk: biller failures pass straight through to partners and can affect partner retention."
p1.font.size = Pt(11)
p1.font.color.rgb = COLOR_TEXT_MAIN

add_card(s7, Inches(4.8), Inches(1.6), c_w, c_h)
tb2 = s7.shapes.add_textbox(Inches(5.0), Inches(1.8), c_w - Inches(0.4), c_h - Inches(0.4))
tf2 = tb2.text_frame
tf2.word_wrap = True
p = tf2.paragraphs[0]
p.text = "Customer Experience"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = COLOR_PRIMARY
p.space_after = Pt(8)

p2 = tf2.add_paragraph()
p2.text = f"• {metrics['unique_customers']:,} unique customer MSISDNs analyzed.\n• 32,150 failed transactions, each a failed purchase attempt for a customer.\n• Repeated failures risk customers abandoning the purchase or switching apps.\n• Clear status messages and safe retries reduce that friction."
p2.font.size = Pt(11)
p2.font.color.rgb = COLOR_TEXT_MAIN

add_card(s7, Inches(8.8), Inches(1.6), c_w, c_h, bg_color=RGBColor(240, 253, 250), border_color=COLOR_ACCENT_TEAL)
tb3 = s7.shapes.add_textbox(Inches(9.0), Inches(1.8), c_w - Inches(0.4), c_h - Inches(0.4))
tf3 = tb3.text_frame
tf3.word_wrap = True
p = tf3.paragraphs[0]
p.text = "Recovery Potential (illustrative)"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = COLOR_ACCENT_TEAL
p.space_after = Pt(8)

p3 = tf3.add_paragraph()
p3.text = "• Every 10% of failures recovered is about 3,200 transactions per 3-day window.\n• The two worst billers account for 19,585 failures, the largest single lever.\n• Revenue impact cannot be estimated: the data has no transaction value."
p3.font.size = Pt(11)
p3.font.color.rgb = COLOR_TEXT_MAIN

# ==============================================================================
# SLIDE 8: STRATEGIC RECOMMENDATIONS & OPERATIONAL ROADMAP
# ==============================================================================
s8 = prs.slides.add_slide(blank_layout)
add_header(s8, "Actionable Strategy: 4-Pillar Operational Optimization")

strat_cards = [
    ("1. Backup Biller & Health-Based Routing",
     "Onboard a second supplier for the XDF and TNP product lines (none exists in this data). Then switch traffic automatically when a biller's failure rate passes 10% in a 5-minute window.",
     COLOR_PRIMARY),
    ("2. Automated Retry with Exponential Backoff", 
     "For transient errors on top SKUs (TNP23, STU15), retry up to 3 times with backoff before returning a final 'FAILED' status to the customer.",
     COLOR_ACCENT_TEAL),
    ("3. Biller Performance Targets",
     "Agree a success-rate target (e.g. 95%) with Biller_29 and Biller_26, review it weekly, and add deposit-balance alerts.",
     COLOR_DANGER),
    ("4. Peak-Hour Monitoring",
     "Watch failure rate and latency closely between 16:30 and 20:00 WIB. The data shows no capacity problem yet, so any scaling should follow latency data.",
     RGBColor(217, 119, 6))
]

for i, (title, desc, col) in enumerate(strat_cards):
    row = i // 2
    col_idx = i % 2
    c_left = Inches(0.8) + col_idx * Inches(5.9 + 0.3)
    c_top = Inches(1.6) + row * Inches(2.6 + 0.2)
    
    add_card(s8, c_left, c_top, Inches(5.9), Inches(2.6))
    tb = s8.shapes.add_textbox(c_left + Inches(0.2), c_top + Inches(0.2), Inches(5.5), Inches(2.2))
    tf = tb.text_frame
    tf.word_wrap = True
    
    p0 = tf.paragraphs[0]
    p0.text = title
    p0.font.size = Pt(13)
    p0.font.bold = True
    p0.font.color.rgb = col
    p0.space_after = Pt(6)
    
    p1 = tf.add_paragraph()
    p1.text = desc
    p1.font.size = Pt(10.5)
    p1.font.color.rgb = COLOR_TEXT_MAIN

# ==============================================================================
# SLIDE 9: PORTFOLIO HIGHLIGHTS & TECH STACK
# ==============================================================================
s9 = prs.slides.add_slide(blank_layout)
add_header(s9, "Technical Implementation & Portfolio Deliverables")

tech_boxes = [
    ("SQL & Database Engineering", 
     ["Indexed SQLite schema (`idx_trans_date`, `idx_status`, `idx_biller_name`).",
      "Engineered analytical SQL views (`v_daily_biller_summary`, `v_hourly_traffic`).",
      "Advanced queries: CTEs, Window Functions (`RANK()`, `SUM() OVER`), and Pareto cumulative distributions.",
      "Product codes normalised to upper case (318 raw codes, 294 real SKUs)."],
     COLOR_PRIMARY),
    ("Python Data Analytics Pipeline",
     ["Pandas for high-volume data cleaning, type coercion & aggregation.",
      "Anonymized PII with deterministic hashing & mapping verification.",
      "Automated visualization generator (Seaborn & Matplotlib at 300 DPI).",
      "Automated PowerPoint reporting engine via `python-pptx`."],
     COLOR_ACCENT_TEAL),
    ("Business Impact & BI Modeling",
     ["Pareto product volume concentration.",
      "Biller failure concentration analysis.",
      "24-hour load and failure profile.",
      "Recommendations limited to what the data can support."],
     COLOR_SUCCESS)
]

c_w9 = Inches(3.7)
c_h9 = Inches(5.3)

for i, (title, items, col) in enumerate(tech_boxes):
    c_left = Inches(0.8) + i * Inches(3.7 + 0.3)
    add_card(s9, c_left, Inches(1.6), c_w9, c_h9)
    
    tb = s9.shapes.add_textbox(c_left + Inches(0.2), Inches(1.8), c_w9 - Inches(0.4), c_h9 - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = col
    p.space_after = Pt(10)
    
    for it in items:
        p_it = tf.add_paragraph()
        p_it.text = f"• {it}"
        p_it.font.size = Pt(10.5)
        p_it.font.color.rgb = COLOR_TEXT_MAIN
        p_it.space_after = Pt(6)

# ==============================================================================
# SLIDE 10: CLOSING / CONTACT SLIDE
# ==============================================================================
s10 = prs.slides.add_slide(blank_layout)

bg10 = s10.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
bg10.fill.solid()
bg10.fill.fore_color.rgb = COLOR_BG_DARK
bg10.line.fill.background()

tb10 = s10.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(11.333), Inches(4.0))
tf10 = tb10.text_frame
tf10.word_wrap = True

p0 = tf10.paragraphs[0]
p0.text = "PORTFOLIO SUMMARY"
p0.font.size = Pt(14)
p0.font.bold = True
p0.font.color.rgb = RGBColor(56, 189, 248)
p0.space_after = Pt(12)

p1 = tf10.add_paragraph()
p1.text = "Thank You / Q&A"
p1.font.size = Pt(36)
p1.font.bold = True
p1.font.color.rgb = COLOR_WHITE
p1.space_after = Pt(14)

p2 = tf10.add_paragraph()
p2.text = "This end-to-end project demonstrates business problem framing, SQL data modeling, exploratory data analysis, root-cause diagnostics, and executive presentation design."
p2.font.size = Pt(16)
p2.font.color.rgb = RGBColor(203, 213, 225)
p2.space_after = Pt(25)

p3 = tf10.add_paragraph()
p3.text = "Candidate: Novaldi Ramadhan Waluyo | Email: novaldiramadhan28@gmail.com | GitHub: github.com/Achimedes28"
p3.font.size = Pt(13)
p3.font.color.rgb = RGBColor(148, 163, 184)

output_pptx = os.path.join(output_dir, 'Fintech_PPOB_Transaction_Operations_Analytics.pptx')
prs.save(output_pptx)
print(f'PowerPoint presentation saved successfully: {output_pptx}')
